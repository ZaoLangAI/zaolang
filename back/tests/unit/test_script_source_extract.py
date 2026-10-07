"""Local + vision extractors behind `POST /v1/scripts/extract`."""

from __future__ import annotations

import io

import pytest
from PIL import Image
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.errors import ValidationFailed
from app.domain.script_writing.extract import (
    MAX_SOURCE_BYTES,
    MAX_SOURCE_TEXT_LEN,
    extract_script_source,
)
from app.models import AgentRun, User
from app.models.enums import AgentRunStatus
from tests.fake_llm_gateway import FAKE_EXTRACTED_SCRIPT


def _png_bytes() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), (30, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def _docx_bytes(text: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "角色"
    table.rows[0].cells[1].text = "林夏"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_extracts_utf8_txt(db: Session, author: User) -> None:
    result = extract_script_source(
        db,
        user_id=author.id,
        filename="idea.txt",
        payload="深夜便利店的秘密\n第二行".encode(),
        mime_type="text/plain",
    )
    assert "深夜便利店的秘密" in result.text
    assert result.truncated is False
    assert result.char_count == len(result.text)


def test_extracts_gb18030_txt(db: Session, author: User) -> None:
    result = extract_script_source(
        db,
        user_id=author.id,
        filename="idea.txt",
        payload="深夜便利店的秘密".encode("gb18030"),
        mime_type="text/plain",
    )
    assert result.text == "深夜便利店的秘密"


def test_extracts_docx_paragraphs_and_tables(db: Session, author: User) -> None:
    result = extract_script_source(
        db,
        user_id=author.id,
        filename="play.docx",
        payload=_docx_bytes("第一场 便利店"),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "第一场 便利店" in result.text
    assert "林夏" in result.text


def test_misnamed_docx_as_doc_still_parses(db: Session, author: User) -> None:
    result = extract_script_source(
        db,
        user_id=author.id,
        filename="play.doc",
        payload=_docx_bytes("误用 doc 后缀的 docx"),
        mime_type="application/msword",
    )
    assert "误用 doc 后缀的 docx" in result.text


def test_garbage_doc_asks_for_conversion(db: Session, author: User) -> None:
    with pytest.raises(ValidationFailed, match="另存为"):
        extract_script_source(
            db,
            user_id=author.id,
            filename="old.doc",
            payload=b"not-a-word-file",
            mime_type="application/msword",
        )


def test_empty_file_is_rejected(db: Session, author: User) -> None:
    with pytest.raises(ValidationFailed, match="空"):
        extract_script_source(
            db, user_id=author.id, filename="empty.txt", payload=b"", mime_type="text/plain"
        )


def test_oversize_payload_is_rejected(db: Session, author: User) -> None:
    with pytest.raises(ValidationFailed, match="8MB"):
        extract_script_source(
            db,
            user_id=author.id,
            filename="big.txt",
            payload=b"x" * (MAX_SOURCE_BYTES + 1),
            mime_type="text/plain",
        )


def test_unsupported_type_is_rejected(db: Session, author: User) -> None:
    with pytest.raises(ValidationFailed, match="仅支持"):
        extract_script_source(
            db,
            user_id=author.id,
            filename="notes.pdf",
            payload=b"%PDF-1.4",
            mime_type="application/pdf",
        )


def test_whitespace_only_txt_is_empty(db: Session, author: User) -> None:
    with pytest.raises(ValidationFailed, match="可用文本"):
        extract_script_source(
            db,
            user_id=author.id,
            filename="blank.txt",
            payload=b"  \n\n  ",
            mime_type="text/plain",
        )


def test_long_txt_is_truncated(db: Session, author: User) -> None:
    payload = ("字" * (MAX_SOURCE_TEXT_LEN + 50)).encode()
    result = extract_script_source(
        db, user_id=author.id, filename="long.txt", payload=payload, mime_type="text/plain"
    )
    assert result.truncated is True
    assert result.char_count == MAX_SOURCE_TEXT_LEN


def test_png_uses_fake_gateway_ocr(db: Session, author: User) -> None:
    result = extract_script_source(
        db,
        user_id=author.id,
        filename="page.png",
        payload=_png_bytes(),
        mime_type="image/png",
    )
    assert result.text == FAKE_EXTRACTED_SCRIPT
    run = db.scalars(select(AgentRun).where(AgentRun.user_id == author.id)).one()
    assert run.prompt_slot == "script_extract"
    assert run.status == AgentRunStatus.SUCCEEDED


def _seed_extract_endpoints(db: Session, *, vision: bool) -> None:
    """Copy's default agent pinned to a text-only endpoint; optionally an
    image-capable backup elsewhere in the pool."""
    from app.domain.agent_skills import service as agent_skills_service
    from app.platform_config import service as config_service

    endpoints: dict = {
        "text-ep": {
            "name": "文本端点",
            "base_url": "https://text.invalid",
            "api_key": "k",
            "kind": "general",
            "model": "text-model",
            "role": "primary",
            "input_modalities": ["text"],
        }
    }
    if vision:
        endpoints["vision-ep"] = {
            "name": "识图端点",
            "base_url": "https://vision.invalid",
            "api_key": "k",
            "kind": "general",
            "model": "vision-model",
            "role": "backup",
            "input_modalities": ["text", "image"],
        }
    config_service.set_value(
        db, "llm_providers", {"endpoints": endpoints}, actor_user_id=None, note="test"
    )
    agent_skills_service.ensure_default_nodes(db)
    agent_skills_service.ensure_default_profiles(db)
    copy_default = agent_skills_service.default_profile(db, "copy")
    assert copy_default is not None
    agent_skills_service.update_profile(db, copy_default.id, default_endpoint_id="text-ep")


@pytest.mark.real_gateway_seams
def test_image_extract_skips_the_text_only_copy_binding(
    db: Session, author: User, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.llm import client as llm_client

    _seed_extract_endpoints(db, vision=True)
    served: list[str] = []

    def _gateway(**kw):  # type: ignore[no-untyped-def]
        served.append(kw["model"])
        return {
            "model": kw["model"],
            "choices": [{"message": {"content": "第一场 便利店"}, "finish_reason": "stop"}],
            "usage": {"prompt_tokens": 1, "completion_tokens": 1},
        }

    monkeypatch.setattr(llm_client, "_call_gateway", _gateway)

    result = extract_script_source(
        db, user_id=author.id, filename="page.png", payload=_png_bytes(), mime_type="image/png"
    )

    assert result.text == "第一场 便利店"
    assert served == ["vision-model"]
    run = db.scalars(select(AgentRun).where(AgentRun.user_id == author.id)).one()
    assert run.endpoint_id == "vision-ep"


@pytest.mark.real_gateway_seams
def test_image_extract_without_an_image_endpoint_says_so(
    db: Session, author: User, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.domain.errors import NoCapableEndpoint
    from app.llm import client as llm_client

    _seed_extract_endpoints(db, vision=False)

    def _gateway(**_kw):  # type: ignore[no-untyped-def]
        raise AssertionError("a text-only endpoint must not receive the image")

    monkeypatch.setattr(llm_client, "_call_gateway", _gateway)

    with pytest.raises(NoCapableEndpoint, match="未配置支持图像输入的端点"):
        extract_script_source(
            db, user_id=author.id, filename="page.png", payload=_png_bytes(), mime_type="image/png"
        )
