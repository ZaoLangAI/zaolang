"""Extract usable script text from an uploaded source file.

txt / docx / (best-effort) doc are parsed locally. jpg / png go through the
copy agent's LLM binding as a one-shot vision call so a photographed
screenplay can become the same `idea` string `prepare_new_script` already
accepts. The call only runs on endpoints that declare `"image"` input: a
copy agent pinned to text-only endpoints falls back to the capable shared
pool, and with none configured the upload fails with `NoCapableEndpoint`
("未配置支持图像输入的端点") instead of reaching a model that cannot see
it. The original bytes are not persisted — only the extracted text is,
later, on `DramaEpisode.source_idea`.
"""

from __future__ import annotations

import base64
import io
import re
import zipfile
from dataclasses import dataclass

from PIL import Image, UnidentifiedImageError
from sqlalchemy.orm import Session

from app.agents.base import GATEWAY_MODE, _record_agent_run, effective_binding
from app.agents.slots import SCRIPT_EXTRACT_SLOT
from app.domain.agent_skills import service as agent_skills_service
from app.domain.errors import ProviderTemporaryFailure, ValidationFailed
from app.llm import client as llm_client
from app.models.enums import AgentName, AgentRunStatus

MAX_SOURCE_BYTES = 8 * 1024 * 1024
MAX_SOURCE_TEXT_LEN = 20_000
# Distinctive marker so `tests/fake_llm_gateway.py` can dispatch this call
# without sniffing image bytes.
SCRIPT_EXTRACT_MARKER = "script_source_extract"
SCRIPT_EXTRACT_SYSTEM_PROMPT = (
    f"[{SCRIPT_EXTRACT_MARKER}] 你是造浪平台的剧本原文提取助手。"
    "从用户给出的图片中识别全部可读文字。"
    "若内容是剧本或分场大纲，保留场次标题、角色名与台词换行。"
    "只输出提取的原文，不要解释、不要翻译、不要补写。"
)

_DOC_CONVERT_HINT = "旧版 .doc 无法解析，请另存为 .docx 或 .txt。"
_EMPTY_HINT = "未能从文件中提取到可用文本。"
_TYPE_HINT = "仅支持 txt、doc、docx、jpg、png。"

_EXT_MIME: dict[str, str] = {
    ".txt": "text/plain",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
}
_MIME_ALIASES: dict[str, str] = {
    "text/plain": "text/plain",
    "application/msword": "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
    "image/jpeg": "image/jpeg",
    "image/jpg": "image/jpeg",
    "image/png": "image/png",
}
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@dataclass(frozen=True, slots=True)
class ExtractedScriptSource:
    filename: str
    text: str
    char_count: int
    truncated: bool


def extract_script_source(
    session: Session,
    *,
    user_id: str,
    filename: str,
    payload: bytes,
    mime_type: str = "",
) -> ExtractedScriptSource:
    if not payload:
        raise ValidationFailed("文件是空的。")
    if len(payload) > MAX_SOURCE_BYTES:
        raise ValidationFailed("文件超过 8MB 上限。")

    kind = _resolved_kind(filename, mime_type)
    if kind == "txt":
        text = _extract_txt(payload)
    elif kind == "docx":
        text = _extract_docx(payload)
    elif kind == "doc":
        text = _extract_doc(payload)
    elif kind in {"jpg", "png"}:
        text = _extract_image(session, payload, kind=kind, filename=filename, user_id=user_id)
    else:
        raise ValidationFailed(_TYPE_HINT)

    normalized = _normalize_extracted(text)
    if not normalized:
        raise ValidationFailed(_EMPTY_HINT)
    truncated = len(normalized) > MAX_SOURCE_TEXT_LEN
    clipped = normalized[:MAX_SOURCE_TEXT_LEN]
    safe_name = (filename or "script").strip()[:255] or "script"
    return ExtractedScriptSource(
        filename=safe_name,
        text=clipped,
        char_count=len(clipped),
        truncated=truncated,
    )


def _resolved_kind(filename: str, mime_type: str) -> str:
    ext = ""
    if "." in filename:
        ext = "." + filename.rsplit(".", 1)[-1].lower()
        if ext == ".jpeg":
            ext = ".jpg"
    declared = _MIME_ALIASES.get((mime_type or "").split(";")[0].strip().lower(), "")
    if ext in _EXT_MIME:
        return ext.lstrip(".")
    if declared == "text/plain":
        return "txt"
    if declared == "application/msword":
        return "doc"
    if declared.endswith("wordprocessingml.document"):
        return "docx"
    if declared == "image/jpeg":
        return "jpg"
    if declared == "image/png":
        return "png"
    raise ValidationFailed(_TYPE_HINT)


def _extract_txt(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


def _extract_docx(payload: bytes) -> str:
    if not _looks_like_zip(payload):
        raise ValidationFailed("无法解析该 docx 文件。")
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - pinned in requirements
        raise ValidationFailed("服务器缺少 docx 解析组件。") from exc
    try:
        document = Document(io.BytesIO(payload))
    except Exception as exc:
        raise ValidationFailed("无法解析该 docx 文件。") from exc
    parts: list[str] = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _extract_doc(payload: bytes) -> str:
    if _looks_like_zip(payload):
        return _extract_docx(payload)
    if not payload.startswith(_OLE_MAGIC):
        raise ValidationFailed(_DOC_CONVERT_HINT)
    try:
        import olefile
    except ImportError as exc:  # pragma: no cover - pinned in requirements
        raise ValidationFailed(_DOC_CONVERT_HINT) from exc
    bio = io.BytesIO(payload)
    if not olefile.isOleFile(bio):
        raise ValidationFailed(_DOC_CONVERT_HINT)
    bio.seek(0)
    ole = olefile.OleFileIO(bio)
    try:
        if not ole.exists("WordDocument"):
            raise ValidationFailed(_DOC_CONVERT_HINT)
        streams = [ole.openstream("WordDocument").read()]
        for name in ("1Table", "0Table"):
            if ole.exists(name):
                streams.append(ole.openstream(name).read())
    except ValidationFailed:
        raise
    except Exception as exc:
        raise ValidationFailed(_DOC_CONVERT_HINT) from exc
    finally:
        ole.close()
    text = _readable_strings_from_ole_streams(streams)
    if len("".join(text.split())) < 20:
        raise ValidationFailed(_DOC_CONVERT_HINT)
    return text


def _extract_image(
    session: Session,
    payload: bytes,
    *,
    kind: str,
    filename: str,
    user_id: str,
) -> str:
    mime = "image/png" if kind == "png" else "image/jpeg"
    try:
        with Image.open(io.BytesIO(payload)) as image:
            image.verify()
        with Image.open(io.BytesIO(payload)) as image:
            fmt = (image.format or "").upper()
    except (UnidentifiedImageError, OSError) as exc:
        raise ValidationFailed("无法解析该图片文件。") from exc
    if fmt not in {"JPEG", "PNG"}:
        raise ValidationFailed("仅支持 jpg、png 图片。")

    data_url = f"data:{mime};base64,{base64.b64encode(payload).decode('ascii')}"
    agent_id = agent_skills_service.resolve_copy_agent_id(session)
    resolved = agent_skills_service.resolve_prompt(
        session,
        AgentName.COPY.value,
        SCRIPT_EXTRACT_SYSTEM_PROMPT,
        agent_id=agent_id,
        slot=SCRIPT_EXTRACT_SLOT,
    )
    binding = effective_binding(
        session, AgentName.COPY.value, resolved.profile, required_modalities={"image"}
    )
    try:
        result = llm_client.complete(
            session=session,
            agent_name=AgentName.COPY.value,
            model=binding.model or "",
            messages=[
                {"role": "system", "content": SCRIPT_EXTRACT_SYSTEM_PROMPT},
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": (
                                "请提取这张图片中的全部可读文字。"
                                "若是剧本，保留场次、角色名与台词换行。只输出原文。"
                            ),
                        },
                        {"type": "image_url", "image_url": {"url": data_url}},
                    ],
                },
            ],
            max_tokens=max(2048, binding.max_tokens),
            temperature=0.0,
            expect_json=False,
            reasoning_model=binding.reasoning_model,
            preferred_endpoint_ids=binding.preferred_endpoint_ids,
        )
    except ProviderTemporaryFailure:
        raise

    text = (result.response.text or "").strip()
    parse_failed = not text
    _record_agent_run(
        session,
        agent_name=AgentName.COPY.value,
        profile_id=resolved.profile.id if resolved.profile is not None else None,
        slot=SCRIPT_EXTRACT_SLOT,
        job_id=None,
        user_id=user_id,
        mode=GATEWAY_MODE,
        model=result.response.model or binding.model or None,
        status=AgentRunStatus.FAILED if parse_failed else AgentRunStatus.SUCCEEDED,
        degraded=parse_failed,
        degrade_reason="empty_extract" if parse_failed else None,
        prompt_tokens=result.response.prompt_tokens,
        completion_tokens=result.response.completion_tokens,
        latency_ms=result.latency_ms,
        endpoint_id=result.endpoint_id,
        input_json={
            "system_prompt": SCRIPT_EXTRACT_SYSTEM_PROMPT,
            "user_prompt": "[image]",
            "filename": filename,
            "mime_type": mime,
        },
        output_json={"text": text},
        thinking_text=result.thinking,
    )
    if parse_failed:
        raise ValidationFailed(_EMPTY_HINT)
    return text


def _looks_like_zip(payload: bytes) -> bool:
    if not payload.startswith(b"PK"):
        return False
    try:
        return zipfile.is_zipfile(io.BytesIO(payload))
    except Exception:
        return False


def _readable_strings_from_ole_streams(streams: list[bytes]) -> str:
    """Best-effort text scrape from a Word 97 `WordDocument` stream."""
    parts: list[str] = []
    utf16_run = re.compile(rb"(?:[\x20-\x7e\n\r\t]\x00){6,}")
    ascii_run = re.compile(rb"[\x20-\x7e]{8,}")
    cjk_run = re.compile(rb"(?:(?:[\x20-\x7e\n\r\t]\x00)|(?:[\x00-\xff][\x4e-\x9f])){6,}")
    for data in streams:
        for match in utf16_run.finditer(data):
            parts.append(match.group().decode("utf-16le", errors="ignore"))
        for match in cjk_run.finditer(data):
            parts.append(match.group().decode("utf-16le", errors="ignore"))
        for match in ascii_run.finditer(data):
            parts.append(match.group().decode("ascii", errors="ignore"))
    cleaned: list[str] = []
    seen: set[str] = set()
    for part in parts:
        text = _normalize_extracted(part)
        if not text or text in seen or not _looks_like_prose(text):
            continue
        seen.add(text)
        cleaned.append(text)
    return "\n".join(cleaned)


def _looks_like_prose(text: str) -> bool:
    compact = "".join(text.split())
    if len(compact) < 8:
        return False
    good = sum(1 for ch in compact if ch.isalnum() or "\u4e00" <= ch <= "\u9fff")
    return good / len(compact) >= 0.55


def _normalize_extracted(text: str) -> str:
    unified = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in unified.split("\n")]
    collapsed: list[str] = []
    blank = False
    for line in lines:
        if not line.strip():
            if collapsed and not blank:
                collapsed.append("")
            blank = True
            continue
        blank = False
        collapsed.append(line.strip())
    return "\n".join(collapsed).strip()
